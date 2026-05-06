
import io
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from .helpers import safe_sheet_name

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FILL = PatternFill("solid", fgColor="D9EAF7")
SECTION_FILL = PatternFill("solid", fgColor="E2F0D9")
DANGER_FILL = PatternFill("solid", fgColor="FDE9E7")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
SUCCESS_FILL = PatternFill("solid", fgColor="E2F0D9")
INFO_FILL = PatternFill("solid", fgColor="D9EAF7")
THIN_BORDER = Border(left=Side(style="thin", color="D9D9D9"), right=Side(style="thin", color="D9D9D9"), top=Side(style="thin", color="D9D9D9"), bottom=Side(style="thin", color="D9D9D9"))

def format_sheet(ws, meta_text="", currency_cols=None, percent_cols=None, integer_cols=None):
    currency_cols = currency_cols or set()
    percent_cols = percent_cols or set()
    integer_cols = integer_cols or set()

    ws.insert_rows(1)
    ws["A1"] = meta_text
    ws["A1"].fill = TITLE_FILL
    ws["A1"].font = Font(bold=True)
    ws["A1"].alignment = Alignment(horizontal="left")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, ws.max_column))

    for cell in ws[2]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center")

    ws.freeze_panes = "A3"

    headers = [ws.cell(row=2, column=i).value for i in range(1, ws.max_column + 1)]
    for idx, header in enumerate(headers, start=1):
        letter = get_column_letter(idx)
        max_len = len(str(header)) if header is not None else 0
        for cell in ws[letter]:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 40)

        if header in currency_cols:
            for r in range(3, ws.max_row + 1):
                ws.cell(r, idx).number_format = '#,##0.00'
        if header in percent_cols:
            for r in range(3, ws.max_row + 1):
                ws.cell(r, idx).number_format = '0.00%'
        if header in integer_cols:
            for r in range(3, ws.max_row + 1):
                ws.cell(r, idx).number_format = '#,##0'

    # Safe autofilter only on the actual header/data range
    if ws.max_row >= 2 and ws.max_column >= 1:
        ws.auto_filter.ref = f"A2:{get_column_letter(ws.max_column)}{ws.max_row}"

    head = {headers[i - 1]: i for i in range(1, len(headers) + 1)}
    if "Need" in head:
        c = get_column_letter(head["Need"])
        ws.conditional_formatting.add(f"{c}3:{c}{ws.max_row}", CellIsRule(operator="greaterThan", formula=["0"], fill=DANGER_FILL))
        ws.conditional_formatting.add(f"{c}3:{c}{ws.max_row}", CellIsRule(operator="equal", formula=["0"], fill=SUCCESS_FILL))
    if "Priority" in head:
        c = get_column_letter(head["Priority"])
        for label, fill in [("High", DANGER_FILL), ("Medium", WARN_FILL), ("Low", INFO_FILL)]:
            ws.conditional_formatting.add(f"{c}3:{c}{ws.max_row}", CellIsRule(operator="equal", formula=[f'"{label}"'], fill=fill))
    if "Status" in head:
        c = get_column_letter(head["Status"])
        for label, fill in [("NEED ORDER", DANGER_FILL), ("ENOUGH", SUCCESS_FILL), ("Below", WARN_FILL), ("Above", SUCCESS_FILL), ("Missing", DANGER_FILL), ("Weak Sales", WARN_FILL), ("Covered", SUCCESS_FILL), ("High", DANGER_FILL), ("Medium", WARN_FILL), ("OK", SUCCESS_FILL), ("Info", INFO_FILL), ("Not Loaded", WARN_FILL)]:
            ws.conditional_formatting.add(f"{c}3:{c}{ws.max_row}", CellIsRule(operator="equal", formula=[f'"{label}"'], fill=fill))
    if "Risk" in head:
        c = get_column_letter(head["Risk"])
        for label, fill in [("High", DANGER_FILL), ("Medium", WARN_FILL), ("Low", SUCCESS_FILL)]:
            ws.conditional_formatting.add(f"{c}3:{c}{ws.max_row}", CellIsRule(operator="equal", formula=[f'"{label}"'], fill=fill))
    if "ABC Class (Qty 70%)" in head:
        c = get_column_letter(head["ABC Class (Qty 70%)"])
        for label, fill in [("A", DANGER_FILL), ("B", WARN_FILL), ("C", INFO_FILL)]:
            ws.conditional_formatting.add(f"{c}3:{c}{ws.max_row}", CellIsRule(operator="equal", formula=[f'"{label}"'], fill=fill))
    if "Loaded" in head:
        c = get_column_letter(head["Loaded"])
        for label, fill in [("Yes", SUCCESS_FILL), ("No", WARN_FILL)]:
            ws.conditional_formatting.add(f"{c}3:{c}{ws.max_row}", CellIsRule(operator="equal", formula=[f'"{label}"'], fill=fill))

def to_excel_bytes(report_sheets, meta):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for name, df in report_sheets.items():
            df.to_excel(writer, sheet_name=safe_sheet_name(name), index=False)

        wb = writer.book
        currency_cols = {"Total Amount","Total Sales","Stock Value","Retail Value","Total Sales Amount (Period)","SaleAmt","RetailValue","Retail Price","Avg Basket","Total Sales Without VAT","Total Sales With VAT","Home Delivery Sales","Wasfaty Sales","Private Label Sales","Cash Sales","Value","NUPCO Price","Commission","DRP","RP","Applied Offer","SReturn","BS"}
        percent_cols = {"Share%","Cum%","Achievement %","Wasfaty %","Home Delivery %","Private Label %","Dead Stock %","PL_Mix","Promoted Sales Mix","NEMix","Branch_Share","Region_Share"}
        integer_cols = {"Total Qty","Stock Qty","Need","Sales Days","Customer Count","Critical Actions","Opportunities","Below Region Categories","Total Sales Qty (Period)","High Risk Items","Staff Count","Low Staff Count (<60 BS)","Rows"}

        for name in report_sheets.keys():
            ws = wb[safe_sheet_name(name)]
            meta_text = f"Branch: {meta['branch']} | Period: {meta['months']} Months | Level: {meta['level']} | Generated: {meta['generated']}"
            format_sheet(ws, meta_text, currency_cols, percent_cols, integer_cols)

            if ws.title == "Executive_Summary":
                for r in range(3, ws.max_row + 1):
                    ws.cell(r, 1).fill = SECTION_FILL
                    ws.cell(r, 1).font = Font(bold=True)

    output.seek(0)
    return output.getvalue()
